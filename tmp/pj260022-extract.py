import json
from pathlib import Path
import openpyxl

source = Path(r'C:/Users/natta/OneDrive - Tomas Tech/IoT Team - Documents/Project - 2026/[PJ260022] SR Jig for Howa/20260121-001-1_Estimate cost for HOWA - QR Reader Jig.xlsx')
w = openpyxl.load_workbook(source, data_only=True, read_only=True)
f = openpyxl.load_workbook(source, data_only=False, read_only=True)
summary = w['001_Summary cost']
def clean(x):
    return x.strip() if isinstance(x, str) else x
def row(sheet, r):
    s=w[sheet]
    result={'sheet':sheet,'row':r}
    for key,col in [('model','D'),('description','E'),('supplier','F'),('brand','G'),('lead_time','H'),('quote_revision','I'),('unit_price','J'),('quantity','K'),('total','L'),('unit','M'),('remark','N')]:
        result[key]=clean(s[f'{col}{r}'].value)
    result['formulas']={c.coordinate:c.value for rr in f[sheet].iter_rows(min_row=r,max_row=r,min_col=4,max_col=14) for c in rr if c.data_type=='f'}
    return result

summary_rows=[9,10,11,14,19,20,21,24,29,30,31,32]
summary_items=[row(summary.title,r) for r in summary_rows]
details={
    'standard':[row('002_STD',r) for r in [10,35,58]],
    'electrical_material':[row('003_COST_EE',r) for r in range(10,15)],
    'electrical_labor':[row('003_COST_EE',r) for r in range(131,136)],
    'mechanical_material':[row('004_COST_ME',r) for r in range(11,32)]+[row('004_COST_ME',34)],
    'mechanical_labor':[row('004_COST_ME',r) for r in range(186,192)],
}
items=[]
def add(source_row, category, summary_row, multiplier=1, quantity=None, unit_price=None, note=''):
    item=dict(source_row)
    item.update(category=category,summary_row=summary_row,quantity=quantity if quantity is not None else float(source_row['quantity'] or 0)*multiplier,unit_price=unit_price if unit_price is not None else float(source_row['unit_price'] or 0),note=note)
    item['total']=round(item['quantity']*item['unit_price'],2)
    items.append(item)
for r,sr in zip(details['standard'],[9,10,11]):
    add(r,'Standard equipment',sr,quantity=summary[f'K{sr}'].value,unit_price=0,note='Customer Provide; quantity from summary; no purchase cost.')
for r in details['electrical_material']:
    add(r,'Electrical material',14,note='Part of Control Box; replaces summary aggregate row14.')
for sr in [19,20,21]:
    add(row(summary.title,sr),'Electrical labor',sr,note='Summary is authoritative; detail quantities differ for drawing and assembly.')
for r in details['mechanical_material']:
    if r['row']==34:
        add(r,'Mechanical safety cost',24,quantity=2,unit_price=970,note='970 per jig fixed amount in L34; J34 is0 and K34 blank. Two jigs from summary K24.')
    else:
        add(r,'Mechanical material',24,multiplier=2,note='BOM per jig multiplied by two jigs from summary K24; replaces summary aggregate row24.')
for sr in [29,30,31,32]:
    add(row(summary.title,sr),'Mechanical labor' if sr!=32 else 'Transportation',sr)
totals={'currency':'THB','direct_cost':summary['L59'].value,'profit_markup_rate':summary['K60'].value,'profit_amount':summary['L60'].value,'overhead_rate':summary['K61'].value,'overhead_amount':summary['L61'].value,'grand_total':summary['L62'].value,'summary_safety_cost':summary['L58'].value,'mechanical_safety_cost_in_direct_cost':1940,'vat':'Not specified in workbook'}
assert sum(i['total'] for i in items)==totals['direct_cost']==60950
assert sum(i['total'] for i in items if i['summary_row']==14)==3950
assert sum(i['total'] for i in items if i['summary_row']==24)==25000
assert totals['direct_cost']*(1+totals['profit_markup_rate']+totals['overhead_rate'])==totals['grand_total']==97520
errors=[{'sheet':s.title,'cell':c.coordinate,'cached':c.value,'formula':f[s.title][c.coordinate].value} for s in w for rr in s for c in rr if c.data_type=='e']
notes=[
 'Use user project PJ260022; summary E5 contains Inquiry, a template header.',
 'Use 001_Summary cost as commercial scope and totals; do not sum every sheet.',
 'Summary19/20 each has quantity1 and cost5000; detail003_COST_EE K131/K132 are0. Import summary values, retain discrepancy note.',
 '003_COST_EE row134 Installation5000 is not referenced by summary and must not add to estimate.',
 '004_COST_ME L34 fixed970 safety cost per jig is included in L35=12500. For two jigs include1940. Do not double count aggregate Sensor Jig Stand25000 when importing its BOM.',
 'Unused 005_COST_Server_PC L60/L122,006_COST_IF L60,007_COST_PG L60 each carry10000 safety placeholders absent from summary. Exclude.',
 '001_Summary cost E46 formula=-ZH500N has cached#NAME? in description only. No quantity/unitprice/lineamount and no reference in financial formulas; exclude junk row.',
 'Profit55% and overhead5% are both calculated on direct cost60950, not gross margin on sales or markup on subtotal+overhead.',
 'Unit for summary31 Installation,Test is Set although detailM189 says Day; preserve summary Set.',
 'Unselected zeros/accommodation, miscellaneous notes and manpower template sheets are not extra cost scope.'
]
out={'source':str(source),'project_code':'PJ260022','project_name':summary['E2'].value,'customer':summary['E3'].value,'document_date':summary['L3'].value.date().isoformat(),'creator':summary['K3'].value,'summary_items':summary_items,'detail_items':details,'recommended_import_items':items,'totals':totals,'summary_total_formulas':{a:f[summary.title][a].value for a in ['L59','L60','L61','L62']},'formula_errors':errors,'notes':notes,'reconciliation':{'item_count':len(items),'direct_cost_sum':sum(i['total'] for i in items),'customer_provided_item_count':3,'electrical_material':3950,'mechanical_material_including_safety':25000,'electrical_labor':15000,'mechanical_labor':15000,'transportation':2000}}
target=Path('tmp/pj260022-extracted.json')
target.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'file':str(target),'reconciliation':out['reconciliation'],'totals':totals,'notes':notes},ensure_ascii=False,indent=2))
